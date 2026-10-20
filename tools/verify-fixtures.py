from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
out=root/'results'
source=root/'A_매출비교'/'01_영업실적.xlsx'
book=load_workbook(source)
ws=book['원본']
ws.cell(3,1).value=ws.cell(2,1).value
book.save(out/'verify-duplicate-within.xlsx')
book=load_workbook(source)
online=load_workbook(root/'A_매출비교'/'02_온라인실적.xlsx',data_only=True)
headers=[c.value for c in online['원본'][1]]
book['원본'].cell(2,1).value=online['원본'].cell(2,headers.index('레코드ID')+1).value
book.save(out/'verify-duplicate-cross.xlsx')
book=load_workbook(source)
headers=[c.value for c in book['원본'][1]]
book['원본'].cell(2,headers.index('채널')+1).value='온라인'
book.save(out/'verify-channel-mismatch.xlsx')
print('Created three independent error fixtures in results/')
