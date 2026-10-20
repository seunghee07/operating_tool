import pathlib, hashlib, json, datetime
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / 'results'
OUT.mkdir(exist_ok=True)
files = sorted(p for pattern in ['A_*', 'B_*', 'C_*', 'D_*', 'E_*'] for folder in ROOT.glob(pattern) if folder.is_dir() for p in folder.rglob('*') if p.is_file())
hashes = {str(p.relative_to(ROOT)).replace('\\', '/'): hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
before = OUT / 'verify-original-before.json'
if not before.exists():
    before.write_text(json.dumps(hashes, ensure_ascii=False, indent=2), encoding='utf-8')
else:
    initial = json.loads(before.read_text(encoding='utf-8'))
    assert initial == hashes, 'Original file hashes changed'
    (OUT / 'verify-original-after.json').write_text(json.dumps(hashes, ensure_ascii=False, indent=2), encoding='utf-8')

aliases = {
 'record_id':['record_id','레코드ID','기록번호'], '기준주':['기준주','주차','주간기준일'],
 '점포':['점포','점포명','영업점'], '브랜드':['브랜드','브랜드명','행사브랜드'],
 '장르':['장르','카테고리','상품군'], '채널':['채널','유통채널','판매채널'],
 '매출_만원':['매출_만원','실적(만원)','매출액_만원'], '목표_만원':['목표_만원','목표매출(만원)','목표액_만원'],
 '비용_만원':['비용_만원','집행비용(만원)','사용비용_만원'], '상태':['상태','진행상태','진행현황'],
 '출처':['출처','근거자료','자료출처'], '기준일':['기준일','확인일','기준일자']}
all_rows=[]
sources=[]
for name, channel in [('01_영업실적.xlsx','영업'),('02_온라인실적.xlsx','온라인'),('03_팝업실적.xlsx','팝업')]:
    book=load_workbook(ROOT/'A_매출비교'/name,data_only=True)
    raw=list(book['원본'].values)
    headers=raw[0]
    indices={k: next(i for i,h in enumerate(headers) if h in names) for k,names in aliases.items()}
    rows=[{k: row[i] for k,i in indices.items()} for row in raw[1:] if any(v is not None for v in row)]
    assert len(rows)==160
    assert all(r['채널']==channel for r in rows)
    assert len({r['record_id'] for r in rows})==len(rows)
    sources.append({'filename':name,'count':len(rows),'sales':sum(r['매출_만원'] for r in rows),'target':sum(r['목표_만원'] for r in rows),'cost':sum(r['비용_만원'] for r in rows),'mapping':{k:headers[i] for k,i in indices.items()}})
    all_rows.extend(rows)
assert len({r['record_id'] for r in all_rows})==480
report_book=load_workbook(ROOT/'A_매출비교'/'04_보고양식.xlsx',data_only=True)
report_headers=list(next(report_book.active.values))
def serial(v):
    if isinstance(v,(datetime.date,datetime.datetime)): return v.strftime('%Y-%m-%d')
    raise TypeError(str(type(v)))
result={'sources':sources,'count':len(all_rows),'sales':sum(r['매출_만원'] for r in all_rows),'target':sum(r['목표_만원'] for r in all_rows),'cost':sum(r['비용_만원'] for r in all_rows),'reportHeaders':report_headers,'zeroRows':[r for r in all_rows if r['목표_만원']==0],'rows':all_rows,'originalFiles':len(hashes)}
(OUT/'verify-expected.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=serial),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ['rows','zeroRows']},ensure_ascii=False))
print('zero target rows:',len(result['zeroRows']))
